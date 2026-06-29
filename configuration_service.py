from __future__ import annotations
import json, uuid
from datetime import datetime, timezone
from io import BytesIO
from openpyxl import load_workbook
from sqlalchemy.orm import Session
from DataDictionaryAdminApp.config import get_settings
from DataDictionaryAdminApp.core.database import get_session_factory
from DataDictionaryAdminApp.model.configuration_tables import AttributePortfolioScope, AttributeBusinessRule, PortfolioReference, PromptReference, UiDisplayConfig
from DataDictionaryAdminApp.model.master_dictionary import MasterDictionary
from DataDictionaryAdminApp.model.audit import AuditLog

SECTOR_MAP={"banks":("FI","Banks"),"bank":("FI","Banks"),"insurance":("FI","Insurance"),"corporates":("Corporate","Corporate"),"corporate":("Corporate","Corporate"),"downstream":("UKC","UKC"),"ukc":("UKC","UKC"),"snp":("SnP","SnP"),"s&p":("SnP","SnP")}
FLAG_TO_SECTOR={"required_by_banks":("FI","Banks"),"required_by_insurance":("FI","Insurance"),"required_by_corporates":("Corporate","Corporate"),"required_by_downstream":("UKC","UKC")}

class ConfigurationService:
    def __init__(self): self.settings=get_settings()
    def _session(self): return get_session_factory()()
    @staticmethod
    def _as_dict(x): return {c.name:getattr(x,c.name) for c in x.__table__.columns}
    @staticmethod
    def _actor(x): return x or "sysuser"
    def _audit(self, db, prj_id, table, action, user, old=None, new=None):
        db.add(AuditLog(batch_id=str(uuid.uuid4()),prj_id=prj_id,table_name=table,action_type=action,source_module="PART3_CONFIGURATION",old_value=json.dumps(old,default=str) if old else None,new_value=json.dumps(new,default=str) if new else None,changed_by=self._actor(user)))
    def bootstrap_portfolios(self,user_id="sysuser"):
        if not self.settings.enable_db:return {"status":"SIMULATED_SUCCESS","records_processed":5}
        seeds=[("FI","Banks","Financial Institutes - Banks"),("FI","Insurance","Financial Institutes - Insurance"),("Corporate","Corporate","Global Corporates"),("UKC","UKC","UK Corporate"),("SnP","SnP","Standard & Poor's")]
        with self._session() as db:
            for p,s,r in seeds:
                if not db.query(PortfolioReference).filter_by(port_name=p,sector_name=s,sub_sector=None).first(): db.add(PortfolioReference(port_name=p,sector_name=s,remark=r,created_by=self._actor(user_id),updated_by=self._actor(user_id)))
            db.commit()
        return {"status":"SUCCESS","records_processed":len(seeds)}
    def _portfolio(self,db,sector):
        mapped=SECTOR_MAP.get((sector or "").strip().lower())
        if not mapped: raise ValueError(f"Unsupported sector: {sector}")
        x=db.query(PortfolioReference).filter_by(port_name=mapped[0],sector_name=mapped[1],is_active=True).first()
        if not x: raise ValueError(f"Portfolio reference missing for sector: {sector}. Run Bootstrap Portfolios.")
        return x
    def sync_scopes(self,prj_id,user_id="sysuser"):
        if not self.settings.enable_db:return []
        with self._session() as db:
            master=db.query(MasterDictionary).filter_by(prj_id=prj_id).first()
            if not master: raise ValueError(f"PRJ ID {prj_id} is not present in Master Dictionary.")
            refs={(x.port_name,x.sector_name):x for x in db.query(PortfolioReference).filter_by(is_active=True).all()}
            existing={x.port_ref_id:x for x in db.query(AttributePortfolioScope).filter_by(prj_id=prj_id).all()}
            wanted={refs[p].port_ref_id for f,p in FLAG_TO_SECTOR.items() if bool(getattr(master,f)) and p in refs}
            for port_id in wanted:
                x=existing.get(port_id)
                if x: x.is_active=True;x.is_deleted=False;x.description=master.prj_attribute_description;x.updated_by=self._actor(user_id)
                else: db.add(AttributePortfolioScope(prj_id=prj_id,port_ref_id=port_id,description=master.prj_attribute_description,created_by=self._actor(user_id),updated_by=self._actor(user_id)))
            for port_id,x in existing.items():
                if port_id not in wanted: x.is_active=False;x.is_deleted=False;x.updated_by=self._actor(user_id)
            db.commit()
            return [self._as_dict(x) for x in db.query(AttributePortfolioScope).filter_by(prj_id=prj_id).all()]
    def _scope(self,db,prj_id,sector,user="sysuser"):
        if not db.query(MasterDictionary).filter_by(prj_id=prj_id,is_deleted=False).first(): raise ValueError(f"PRJ ID {prj_id} is not present in Master Dictionary.")
        self.sync_scopes(prj_id,user)
        ref=self._portfolio(db,sector); x=db.query(AttributePortfolioScope).filter_by(prj_id=prj_id,port_ref_id=ref.port_ref_id).first()
        if not x or not x.is_active: raise ValueError(f"No active portfolio scope for {prj_id}/{sector}. Set the matching Required by flag in Master Dictionary.")
        return x,ref
    def _soft_delete(self,model,key,value,user):
        with self._session() as db:
            x=db.query(model).filter(getattr(model,key)==value).first()
            if not x: raise ValueError("Record not found")
            old=self._as_dict(x);x.is_active=False;x.is_deleted=True;x.deleted_at=datetime.now(timezone.utc);x.deleted_by=self._actor(user);x.updated_by=self._actor(user)
            self._audit(db,getattr(x,"prj_id","CONFIG"),model.__tablename__,"SOFT_DELETE",user,old,self._as_dict(x));db.commit();return self._as_dict(x)
    def reactivate(self,model,key,value,user):
        with self._session() as db:
            x=db.query(model).filter(getattr(model,key)==value).first()
            if not x: raise ValueError("Record not found")
            old=self._as_dict(x);x.is_active=True;x.is_deleted=False;x.deleted_at=None;x.deleted_by=None;x.updated_by=self._actor(user)
            self._audit(db,getattr(x,"prj_id","CONFIG"),model.__tablename__,"REACTIVATE",user,old,self._as_dict(x));db.commit();return self._as_dict(x)
    def list_displays(self,active_only=True):
        with self._session() as db:
            q=db.query(UiDisplayConfig)
            if active_only:q=q.filter_by(is_active=True,is_deleted=False)
            return [self._as_dict(x) for x in q.order_by(UiDisplayConfig.section,UiDisplayConfig.display_order).all()]
    def save_display(self,d,user_id="sysuser"):
        with self._session() as db:
            scope,_=self._scope(db,d["prj_id"],d["sector"],user_id); x=db.query(UiDisplayConfig).filter_by(display_id=d.get("display_id")).first() if d.get("display_id") else None
            vals={k:d.get(k) for k in ("display_order","display_name","section","subsection","view_name","description")};vals["display_name"]=vals["display_name"] or d.get("prj_attribute_name") or d["prj_id"]
            old=self._as_dict(x) if x else None
            if not x:x=UiDisplayConfig(scope_id=scope.scope_id,created_by=self._actor(user_id),updated_by=self._actor(user_id),**vals);db.add(x);action="INSERT"
            else:
                x.scope_id=scope.scope_id;x.is_active=True;x.is_deleted=False;x.updated_by=self._actor(user_id)
                for k,v in vals.items():setattr(x,k,v)
                action="UPDATE"
            db.flush();self._audit(db,d["prj_id"],x.__tablename__,action,user_id,old,self._as_dict(x));db.commit();db.refresh(x);return self._as_dict(x)
    def save_rule(self,d,user_id="sysuser"):
        with self._session() as db:
            scope,_=self._scope(db,d["prj_id"],d["sector"],user_id);x=db.query(AttributeBusinessRule).filter_by(scope_id=scope.scope_id).first()
            vals={k:d.get(k) for k in ("source_abbr_name","editable","symbol","mapping_type","mapping_logic","calculation_logic","business_logic")};vals["source_abbr_name"]=vals["source_abbr_name"] or "SNPAR";old=self._as_dict(x) if x else None
            if not x:x=AttributeBusinessRule(scope_id=scope.scope_id,created_by=self._actor(user_id),updated_by=self._actor(user_id),**vals);db.add(x);action="INSERT"
            else:
                for k,v in vals.items():setattr(x,k,v)
                x.is_active=True;x.is_deleted=False;x.updated_by=self._actor(user_id);action="UPDATE"
            db.flush();self._audit(db,d["prj_id"],x.__tablename__,action,user_id,old,self._as_dict(x));db.commit();db.refresh(x);return self._as_dict(x)
    def list_rules(self,active_only=True):
        with self._session() as db:
            q=db.query(AttributeBusinessRule)
            if active_only:q=q.filter_by(is_active=True,is_deleted=False)
            return [self._as_dict(x) for x in q.order_by(AttributeBusinessRule.id).all()]
    def list_prompts(self,active_only=True):
        with self._session() as db:
            q=db.query(PromptReference)
            if active_only:q=q.filter_by(is_active=True,is_deleted=False)
            return [self._as_dict(x) for x in q.order_by(PromptReference.prj_id).all()]
    def save_prompt(self,d,user_id="sysuser"):
        with self._session() as db:
            scope,ref=self._scope(db,d["prj_id"],d["sector"],user_id);x=db.query(PromptReference).filter_by(scope_id=scope.scope_id,prj_id=d["prj_id"],port_ref_id=ref.port_ref_id).first()
            fields=("required_by_scope","attribute_name","section","sub_section","data_type","calculated_or_reported","calculation_logic","segment","subcomponent_total","attribute_description","examples","source_sheet_name")
            vals={k:d.get(k) for k in fields};old=self._as_dict(x) if x else None
            if not x:x=PromptReference(scope_id=scope.scope_id,prj_id=d["prj_id"],port_ref_id=ref.port_ref_id,created_by=self._actor(user_id),updated_by=self._actor(user_id),**vals);db.add(x);action="INSERT"
            else:
                for k,v in vals.items():
                    if v is not None:setattr(x,k,v) # missing Excel columns never erase DB data
                x.is_active=True;x.is_deleted=False;x.updated_by=self._actor(user_id);action="UPDATE"
            db.flush();self._audit(db,d["prj_id"],x.__tablename__,action,user_id,old,self._as_dict(x));db.commit();db.refresh(x);return self._as_dict(x)
    def upload_prompts(self,file_bytes,sheets,user_id="sysuser"):
        wb=load_workbook(BytesIO(file_bytes),data_only=True); selected=sheets or wb.sheetnames;records=[];errors=[];skipped=[]
        aliases={"prjid":"prj_id","prj id":"prj_id","prj attribute":"attribute_name","attribute name":"attribute_name","sub-section":"sub_section","sub section":"sub_section","data type":"data_type","calculated or reported":"calculated_or_reported","calculation logic":"calculation_logic","subcomponent/total":"subcomponent_total","description (proposed one-shot prompting)":"attribute_description","description":"attribute_description"}
        for sheet in selected:
            if sheet not in wb.sheetnames: errors.append({"sheet":sheet,"error":"Worksheet not found"});continue
            ws=wb[sheet];headers={str(ws.cell(1,c).value or "").strip().lower():c for c in range(1,ws.max_column+1)}
            mapped={aliases[h]:c for h,c in headers.items() if h in aliases}
            if "prj_id" not in mapped:errors.append({"sheet":sheet,"error":"Missing PRJID / PRJ ID column"});continue
            for row in range(2,ws.max_row+1):
                prj=str(ws.cell(row,mapped["prj_id"]).value or "").strip()
                if not prj:continue
                with self._session() as db: exists=bool(db.query(MasterDictionary.id).filter_by(prj_id=prj,is_deleted=False).first())
                if not exists: skipped.append({"sheet":sheet,"row":row,"prj_id":prj,"reason":"Not present in Master Dictionary"});continue
                d={field:ws.cell(row,col).value for field,col in mapped.items() if field!="prj_id"}; d.update({"prj_id":prj,"sector":str(ws.cell(row,headers["sector"]).value or "SnP") if "sector" in headers else "SnP","source_sheet_name":sheet})
                try: records.append(self.save_prompt(d,user_id))
                except Exception as exc: errors.append({"sheet":sheet,"row":row,"prj_id":prj,"error":str(exc)})
        return {"processed":len(records),"skipped_not_in_master":skipped,"errors":errors,"records":records}
    def soft_delete_display(self,i,u="sysuser"):return self._soft_delete(UiDisplayConfig,"display_id",i,u)
    def soft_delete_prompt(self,i,u="sysuser"):return self._soft_delete(PromptReference,"prompt_id",i,u)
    def soft_delete_rule(self,i,u="sysuser"):return self._soft_delete(AttributeBusinessRule,"id",i,u)
    def reactivate_display(self,i,u="sysuser"):return self.reactivate(UiDisplayConfig,"display_id",i,u)
    def reactivate_prompt(self,i,u="sysuser"):return self.reactivate(PromptReference,"prompt_id",i,u)
    def reactivate_rule(self,i,u="sysuser"):return self.reactivate(AttributeBusinessRule,"id",i,u)
