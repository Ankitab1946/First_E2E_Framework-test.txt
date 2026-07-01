from fastapi import APIRouter, Depends, HTTPException, Query
from sqlalchemy import select, text
from sqlalchemy.orm import Session
from DataDictionaryAdminApp.core.database import get_db
from DataDictionaryAdminApp.api.schemas_api import PromptUpsert
from DataDictionaryAdminApp.model.entities import ScanningPromptReference
from DataDictionaryAdminApp.service.data_dictionary_service import DataDictionaryService
from DataDictionaryAdminApp.utils.security import current_user

router = APIRouter(prefix='/prompts', tags=['Prompts'])
def serial(row): return {c.name:getattr(row,c.name) for c in row.__table__.columns}

@router.get('')
def list_prompts(prj_id:str|None=None, include_deleted:bool=False, db:Session=Depends(get_db)):
    stmt=select(ScanningPromptReference)
    if not include_deleted: stmt=stmt.where(ScanningPromptReference.is_active == 1)
    if prj_id: stmt=stmt.where(ScanningPromptReference.prj_id.ilike(f'%{prj_id}%'))
    return [serial(x) for x in db.scalars(stmt.order_by(ScanningPromptReference.prompt_id)).all()]

@router.post('')
def create_prompt(payload:PromptUpsert,user:str|None=Query(None),db:Session=Depends(get_db)):
    try: return serial(DataDictionaryService(db).upsert_prompt(payload,current_user(user),source='UI'))
    except ValueError as e: raise HTTPException(400,str(e))

@router.put('/{prompt_id}')
def update_prompt(prompt_id:int,payload:PromptUpsert,user:str|None=Query(None),db:Session=Depends(get_db)):
    try: return serial(DataDictionaryService(db).upsert_prompt(payload,current_user(user),prompt_id,source='UI'))
    except ValueError as e: raise HTTPException(400,str(e))

@router.delete('/{prompt_id}')
def delete_prompt(prompt_id:int,user:str|None=Query(None),db:Session=Depends(get_db)):
    if not DataDictionaryService(db).soft_delete_prompt(prompt_id,current_user(user)): raise HTTPException(404,'Prompt not found')
    return {'status':'soft_deleted','prompt_id':prompt_id}

@router.post('/{prompt_id}/reactivate')
def reactivate_prompt(prompt_id:int,user:str|None=Query(None),db:Session=Depends(get_db)):
    row=db.get(ScanningPromptReference,prompt_id)
    if not row: raise HTTPException(404,'Prompt not found')
    row.is_active=True; row.updated_by=current_user(user); db.commit()
    return serial(row)


@router.get('/download-latest')
def download_latest_prompts(db: Session = Depends(get_db)):
    """Download the active prompt table as an Excel workbook."""
    try:
        from io import BytesIO
        import pandas as pd
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        rows = db.execute(text("""
            SELECT prompt_id, scope_id, prj_id, port_ref_id, required_by_scope,
                   attribute_name, display_order, section, sub_section, data_type,
                   calculated_or_reported, calculation_logic, segment,
                   attribute_description, created_at, updated_at, created_by, updated_by
            FROM dbo.prj_scanning_prompt_reference_test
            WHERE is_active = 1
            ORDER BY prj_id, display_order, prompt_id
        """)).mappings().all()
        columns = ['Prompt ID','Scope ID','PRJ ID','Port Ref ID','Required By Scope','Attribute Name','Display Order','Section','Sub-Section','Data Type','Calculated or Reported','Calculation Logic','Segment','Attribute Description','Created At','Updated At','Created By','Updated By']
        wb=Workbook(); ws=wb.active; ws.title='Prompts'
        ws.append(columns)
        for row in rows: ws.append([row.get(k) for k in ['prompt_id','scope_id','prj_id','port_ref_id','required_by_scope','attribute_name','display_order','section','sub_section','data_type','calculated_or_reported','calculation_logic','segment','attribute_description','created_at','updated_at','created_by','updated_by']])
        fill=PatternFill(fill_type='solid', fgColor='1F4E78'); font=Font(color='FFFFFF', bold=True)
        for cell in ws[1]: cell.fill=fill; cell.font=font; cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=min(max(max(len(str(c.value or '')) for c in col)+2,12),42)
        output=BytesIO(); wb.save(output)
        from fastapi.responses import Response
        return Response(output.getvalue(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition':'attachment; filename=prompt_latest.xlsx'})
    except Exception as exc:
        raise HTTPException(400, f'Unable to generate Prompt Excel: {exc}')
